"""The exported workbook's formulas are evaluated independently of the engine."""

from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from exports.excel.evaluate import WorkbookEvaluator
from exports.excel.workbook import build_workbook
from financial_engine.engine import run_model
from financial_engine.reference.northline import northline_model


def _close(actual, expected: str) -> None:
    assert abs(Decimal(str(actual)) - Decimal(expected)) < Decimal("0.01")


def test_formula_workbook_reproduces_the_engine():
    model = northline_model()
    result = run_model(model)
    workbook = load_workbook(BytesIO(build_workbook(model, result)))
    evaluator = WorkbookEvaluator(workbook)
    _close(evaluator.decimal("Revenue", "F5"), "110000")
    _close(evaluator.decimal("Income", "F14"), "4550")
    _close(evaluator.decimal("Income", "F23"), "7432.5")
    _close(evaluator.decimal("Balance", "F5"), "11152.75")
    _close(evaluator.decimal("Balance", "F25"), "25202.75")
    _close(evaluator.decimal("DCF", "F11"), "7162.5")
    _close(evaluator.decimal("WACC", "B16"), str(result.wacc.wacc))
    _close(evaluator.decimal("DCF", "B18"), str(result.dcf.enterprise_value))
    _close(evaluator.decimal("Summary", "B6"), str(result.dcf.enterprise_value))
    _close(evaluator.decimal("Summary", "B7"), str(result.equity.equity_value))
    _close(evaluator.decimal("Summary", "B8"), str(result.equity.value_per_share))
    assert evaluator.decimal("Checks", "F5") == 0
    assert evaluator.decimal("Checks", "F6") == 0
    assert evaluator.decimal("Checks", "G7") == 0
    _close(evaluator.decimal("Income", "G14"), "5155")
    _close(evaluator.decimal("Scenarios", "C10"), "112000")
    # Base cell of the formula sensitivity grid: zero WACC offset and zero growth offset.
    _close(evaluator.decimal("Sensitivities", "D9"), str(result.dcf.enterprise_value))
    assert evaluator.decimal("DCF", "F12") == 1
    _close(evaluator.decimal("Sensitivities", "E18"), str(result.dcf.cross_check_enterprise_value))
    assert evaluator.cell("Sensitivities", "H7") == "n/m"
    assert evaluator.cell("Summary", "B16") == "end"
    blue = workbook["Assumptions"]["B3"].font.color.rgb
    assert blue.endswith("0000FF")
    assert str(workbook["Revenue"]["F5"].value).startswith("=")
    assert "!" in str(workbook["Income"]["F5"].value)


def test_mid_year_workbook_uses_half_year_exponents():
    model = northline_model()
    model.assumptions.valuation.discount_convention = "mid"
    result = run_model(model, include_scenarios=False, include_sensitivities=False)
    workbook = load_workbook(BytesIO(build_workbook(model, result)))
    evaluator = WorkbookEvaluator(workbook)
    assert evaluator.decimal("DCF", "F12") == Decimal("0.5")
    assert evaluator.decimal("DCF", "J12") == Decimal("4.5")
    _close(evaluator.decimal("DCF", "B18"), str(result.dcf.enterprise_value))
    _close(evaluator.decimal("Sensitivities", "D9"), str(result.dcf.enterprise_value))
