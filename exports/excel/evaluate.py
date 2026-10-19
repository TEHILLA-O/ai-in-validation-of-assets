"""Evaluate the Excel formulas ValuationLab writes.

OpenPyXL does not calculate workbooks. This interpreter uses Decimal. It covers
arithmetic, exponentiation, cell references, SUM, MIN, MAX and IF. It is the
check that the exported formulas reproduce the calculation engine.
"""

from __future__ import annotations

from decimal import Decimal

from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.workbook.workbook import Workbook

from financial_engine.precision import D


class FormulaError(Exception):
    pass


def _number(value) -> Decimal:
    if isinstance(value, Decimal):
        return value
    if isinstance(value, bool):
        return Decimal(int(value))
    if isinstance(value, (int, float)):
        return D(str(value))
    if isinstance(value, str):
        return D(value)
    raise FormulaError(f"Cannot use {value!r} as a number")


class WorkbookEvaluator:
    def __init__(self, workbook: Workbook):
        self.workbook = workbook
        self.cache: dict[tuple[str, str], object] = {}
        self.stack: set[tuple[str, str]] = set()

    def cell(self, sheet: str, coord: str):
        coord = coord.replace("$", "")
        key = (sheet, coord)
        if key in self.cache:
            return self.cache[key]
        if key in self.stack:
            raise FormulaError(f"Circular reference at {sheet}!{coord}")
        self.stack.add(key)
        raw = self.workbook[sheet][coord].value
        if isinstance(raw, str) and raw.startswith("="):
            value = Parser(raw[1:], sheet, self).parse()
        elif raw is None:
            value = Decimal(0)
        else:
            value = raw
        self.stack.remove(key)
        self.cache[key] = value
        return value

    def decimal(self, sheet: str, coord: str) -> Decimal:
        return _number(self.cell(sheet, coord))


class Parser:
    def __init__(self, text: str, sheet: str, evaluator: WorkbookEvaluator):
        self.text = text.replace(" ", "")
        self.n = len(self.text)
        self.sheet = sheet
        self.evaluator = evaluator
        self.i = 0

    def parse(self):
        value = self._comparison()
        if self.i != self.n:
            raise FormulaError(f"Unparsed '{self.text[self.i:]}' in {self.text}")
        return value

    def _comparison(self):
        value = self._additive()
        for op in ("<=", ">=", "<>", "=", "<", ">"):
            if self.text.startswith(op, self.i):
                self.i += len(op)
                right = self._additive()
                if op == "=" and (isinstance(value, str) or isinstance(right, str)):
                    return str(value) == str(right)
                left, right_n = _number(value), _number(right)
                return {
                    "<=": left <= right_n,
                    ">=": left >= right_n,
                    "<>": left != right_n,
                    "=": left == right_n,
                    "<": left < right_n,
                    ">": left > right_n,
                }[op]
        return value

    def _additive(self):
        value = self._term()
        while self.i < self.n and self.text[self.i] in "+-":
            op = self.text[self.i]
            self.i += 1
            right = self._term()
            value = _number(value) + _number(right) if op == "+" else _number(value) - _number(right)
        return value

    def _term(self):
        value = self._power()
        while self.i < self.n and self.text[self.i] in "*/":
            op = self.text[self.i]
            self.i += 1
            right = self._power()
            value = _number(value) * _number(right) if op == "*" else _number(value) / _number(right)
        return value

    def _power(self):
        value = self._unary()
        if self.i < self.n and self.text[self.i] == "^":
            self.i += 1
            exponent = _number(self._unary())
            base = _number(value)
            if exponent == int(exponent) and exponent >= 0:
                return base ** int(exponent)
            return (base.ln() * exponent).exp()
        return value

    def _unary(self):
        if self.i < self.n and self.text[self.i] == "+":
            self.i += 1
            return self._unary()
        if self.i < self.n and self.text[self.i] == "-":
            self.i += 1
            return -_number(self._unary())
        return self._primary()

    def _primary(self):
        if self.i >= self.n:
            raise FormulaError(f"Unexpected end of {self.text}")
        char = self.text[self.i]
        if char == "(":
            self.i += 1
            value = self._comparison()
            self._eat(")")
            return value
        if char == '"':
            end = self.text.find('"', self.i + 1)
            if end < 0:
                raise FormulaError("Unclosed string")
            value = self.text[self.i + 1:end]
            self.i = end + 1
            return value
        if char.isdigit() or char == ".":
            return self._number_token()
        if char == "$":
            return self.evaluator.cell(self.sheet, self._coord())
        if char.isalpha() or char == "'":
            return self._name()
        raise FormulaError(f"Unexpected '{char}' in {self.text}")

    def _number_token(self) -> Decimal:
        start = self.i
        while self.i < self.n and (self.text[self.i].isdigit() or self.text[self.i] == "."):
            self.i += 1
        return D(self.text[start:self.i])

    def _name(self):
        if self.text[self.i] == "'":
            end = self.text.find("'", self.i + 1)
            sheet = self.text[self.i + 1:end]
            self.i = end + 1
            self._eat("!")
            return self.evaluator.cell(sheet, self._coord())
        start = self.i
        while self.i < self.n and (self.text[self.i].isalnum() or self.text[self.i] == "_"):
            self.i += 1
        name = self.text[start:self.i]
        if self.i < self.n and self.text[self.i] == "(":
            return self._function(name)
        if self.i < self.n and self.text[self.i] == "!":
            self.i += 1
            return self.evaluator.cell(name, self._coord())
        self.i = start
        return self.evaluator.cell(self.sheet, self._coord())

    def _function(self, name: str):
        self._eat("(")
        upper = name.upper()
        if upper == "IF":
            condition = self._comparison()
            self._eat(",")
            if condition:
                chosen = self._argument()
                self._eat(",")
                self._skip_argument()
            else:
                self._skip_argument()
                self._eat(",")
                chosen = self._argument()
            self._eat(")")
            return chosen
        args = []
        if not (self.i < self.n and self.text[self.i] == ")"):
            args.append(self._argument())
            while self.i < self.n and self.text[self.i] == ",":
                self.i += 1
                args.append(self._argument())
        self._eat(")")
        if upper == "SUM":
            total = Decimal(0)
            for arg in args:
                if isinstance(arg, list):
                    total += sum((_number(item) for item in arg), Decimal(0))
                else:
                    total += _number(arg)
            return total
        if upper == "MIN":
            return min(_number(arg) for arg in args)
        if upper == "MAX":
            return max(_number(arg) for arg in args)
        raise FormulaError(f"Unsupported function {name}")

    def _skip_argument(self) -> None:
        depth = 0
        while self.i < self.n:
            char = self.text[self.i]
            if char == '"':
                end = self.text.find('"', self.i + 1)
                self.i = self.n if end < 0 else end + 1
                continue
            if char == "(":
                depth += 1
            elif char == ")":
                if depth == 0:
                    return
                depth -= 1
            elif char == "," and depth == 0:
                return
            self.i += 1

    def _argument(self):
        saved = self.i
        is_range = self._is_range()
        self.i = saved
        if is_range:
            return self._range()
        return self._comparison()

    def _is_range(self) -> bool:
        try:
            self._consume_ref()
        except FormulaError:
            return False
        return self.i < self.n and self.text[self.i] == ":"

    def _range(self) -> list:
        sheet, start = self._consume_ref()
        self._eat(":")
        end_sheet, end = self._consume_ref()
        owner = sheet or end_sheet or self.sheet
        c1, r1 = _split(start)
        c2, r2 = _split(end)
        values = []
        for col in range(c1, c2 + 1):
            for row in range(r1, r2 + 1):
                values.append(self.evaluator.cell(owner, f"{get_column_letter(col)}{row}"))
        return values

    def _consume_ref(self) -> tuple[str | None, str]:
        sheet = None
        if self.text[self.i] == "'":
            end = self.text.find("'", self.i + 1)
            sheet = self.text[self.i + 1:end]
            self.i = end + 1
            self._eat("!")
        elif self._sheet_ahead():
            start = self.i
            while self.text[self.i].isalnum() or self.text[self.i] == "_":
                self.i += 1
            sheet = self.text[start:self.i]
            self._eat("!")
        return sheet, self._coord()

    def _sheet_ahead(self) -> bool:
        index = self.i
        if not self.text[index].isalpha():
            return False
        while index < self.n and (self.text[index].isalnum() or self.text[index] == "_"):
            index += 1
        return index < self.n and self.text[index] == "!"

    def _coord(self) -> str:
        if self.i < self.n and self.text[self.i] == "$":
            self.i += 1
        start = self.i
        while self.i < self.n and self.text[self.i].isalpha():
            self.i += 1
        if start == self.i:
            raise FormulaError("Missing column")
        if self.i < self.n and self.text[self.i] == "$":
            self.i += 1
        row = self.i
        while self.i < self.n and self.text[self.i].isdigit():
            self.i += 1
        if row == self.i:
            raise FormulaError("Missing row")
        return self.text[start:self.i].replace("$", "")

    def _eat(self, char: str) -> None:
        if self.i >= self.n or self.text[self.i] != char:
            raise FormulaError(f"Expected {char} in {self.text}")
        self.i += 1


def _split(coord: str) -> tuple[int, int]:
    letters = "".join(char for char in coord if char.isalpha())
    digits = "".join(char for char in coord if char.isdigit())
    return column_index_from_string(letters), int(digits)
