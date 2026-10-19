"""Formula-linked Northline-style workbook.

Historical amounts are inputs. Forecast statements, WACC and the selected DCF
are Excel formulas. Vintage depreciation is laid out in rows so a reader can
see why the charge is not a percentage of closing net book value.

Bull and bear enterprise values are values from the same calculation engine,
labelled as such. Their revenue paths are formulas. A full second formula model
for every custom scenario is not duplicated inside the grid; the base case is
the live workbook.
"""

from __future__ import annotations

from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

from financial_engine.model import CompanyModel, ModelResult

NAVY = "1B3A4B"
BLUE = Font(name="Calibri", size=10, color="0000FF")
BLACK = Font(name="Calibri", size=10, color="000000")
GREEN = Font(name="Calibri", size=10, color="008000")
WHITE = Font(name="Calibri", size=10, color="FFFFFF", bold=True)
HEADER = PatternFill("solid", fgColor=NAVY)
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
HIST = PatternFill("solid", fgColor="D6E3F0")
FORE = PatternFill("solid", fgColor="E5F2E9")
SECTION = PatternFill("solid", fgColor="EEF2F5")
TOTAL = PatternFill("solid", fgColor="E7EEF4")
THIN = Border(
    left=Side(style="thin", color="D9E2EC"),
    right=Side(style="thin", color="D9E2EC"),
    top=Side(style="thin", color="D9E2EC"),
    bottom=Side(style="thin", color="D9E2EC"),
)
MONEY = '#,##0.00;(#,##0.00);"-"'
RATE = "0.00%"


def _num(value: Decimal | None) -> float | None:
    if value is None:
        return None
    return float(value)


class ModelWorkbook:
    def __init__(self, model: CompanyModel, result: ModelResult):
        self.model = model
        self.result = result
        self.years = result.years
        self.wb = Workbook()
        self.col = {year.period_id: get_column_letter(3 + index) for index, year in enumerate(self.years)}
        self.ids = [year.period_id for year in self.years]
        self.forecast_ids = [year.period_id for year in self.years if year.is_forecast]
        self.last_actual = next(year.period_id for year in reversed(self.years) if not year.is_forecast)

    def previous(self, period_id: str) -> str | None:
        index = self.ids.index(period_id)
        return None if index == 0 else self.ids[index - 1]

    def build(self) -> bytes:
        default = self.wb.active
        self.wb.remove(default)
        for name in (
            "Summary",
            "Assumptions",
            "Historical",
            "Revenue",
            "OperatingCosts",
            "Income",
            "Balance",
            "CashFlow",
            "WorkingCapital",
            "Capex",
            "Debt",
            "WACC",
            "DCF",
            "Sensitivities",
            "Scenarios",
            "Checks",
            "Sources",
        ):
            self.wb.create_sheet(name)
        self._assumptions()
        self._revenue()
        self._costs()
        self._capex()
        self._debt()
        self._working_capital()
        self._income()
        self._cash_flow()
        self._balance()
        self._wacc()
        self._dcf()
        self._summary()
        self._historical_links()
        self._checks()
        self._sensitivities()
        self._scenarios()
        self._sources()
        for sheet in self.wb.worksheets:
            sheet.freeze_panes = "C5"
            sheet.page_setup.orientation = "landscape"
            sheet.page_setup.fitToPage = True
            sheet.page_setup.fitToWidth = 1
            sheet.page_setup.fitToHeight = 1
            sheet.sheet_view.showGridLines = False
            sheet.oddHeader.left.text = self.model.name
            sheet.oddFooter.left.text = self.model.currency + " in absolute units"
        self.wb.defined_names.add(DefinedName(name="EnterpriseValue", attr_text="Summary!$B$6"))
        buffer = BytesIO()
        self.wb.save(buffer)
        return buffer.getvalue()

    def _put(self, sheet: str, coord: str, value, kind: str = "calc", fmt: str | None = None) -> None:
        cell = self.wb[sheet][coord]
        cell.value = value
        cell.alignment = Alignment(horizontal="right" if coord[0] != "A" else "left", vertical="center")
        cell.border = THIN
        if kind == "input":
            cell.font = BLUE
            cell.fill = INPUT_FILL
        elif kind == "header":
            cell.font = WHITE
            cell.fill = HEADER
            cell.alignment = Alignment(horizontal="center")
        elif kind == "section":
            cell.font = Font(name="Calibri", size=10, bold=True, color=NAVY)
            cell.fill = SECTION
            cell.alignment = Alignment(horizontal="left")
        elif kind == "cross" or (isinstance(value, str) and value.startswith("=") and "!" in value):
            cell.font = GREEN
        elif isinstance(value, str) and value.startswith("="):
            cell.font = BLACK
        else:
            cell.font = BLACK
        if fmt:
            cell.number_format = fmt
        elif isinstance(value, float) or (isinstance(value, str) and value.startswith("=") and kind != "text"):
            cell.number_format = MONEY

    def _banner(self, sheet: str, title: str) -> None:
        ws = self.wb[sheet]
        ws.sheet_view.showGridLines = False
        ws["A1"] = self.model.name + (" — SYNTHETIC" if self.model.synthetic else "")
        ws["A1"].font = Font(name="Calibri", size=16, bold=True, color=NAVY)
        ws["A2"] = title
        ws["A2"].font = Font(name="Calibri", size=10, italic=True, color="52606D")
        ws.column_dimensions["A"].width = 46
        ws.column_dimensions["B"].width = 18
        ws.row_dimensions[1].height = 24
        for index, year in enumerate(self.years):
            letter = get_column_letter(3 + index)
            ws.column_dimensions[letter].width = 14
            self._put(sheet, f"{letter}3", year.label, "header")
            self._put(sheet, f"{letter}4", "Forecast" if year.is_forecast else "Historical", "header")
            ws[f"{letter}4"].fill = FORE if year.is_forecast else HIST

    def _label(self, sheet: str, row: int, text: str, section: bool = False) -> None:
        self._put(sheet, f"A{row}", text, "section" if section else "calc")
        self.wb[sheet][f"A{row}"].font = Font(name="Calibri", size=10, bold=section, color=NAVY)
        self.wb[sheet][f"A{row}"].alignment = Alignment(horizontal="left")
        if section:
            self.wb[sheet][f"A{row}"].fill = SECTION

    def _assumptions(self) -> None:
        assumptions = self.model.assumptions
        self._banner("Assumptions", "Blue cells are inputs. The forecast sheets read these cells.")
        rows = [
            (3, "Tax rate", assumptions.tax_rate, RATE),
            (4, "Interest rate on opening debt", assumptions.interest_rate, RATE),
            (5, "Interest income rate on opening cash", assumptions.interest_income_rate, RATE),
            (6, "Dividend payout of positive net income", assumptions.dividend_payout, RATE),
            (7, "CapEx as a percentage of revenue", assumptions.capex_pct_revenue, RATE),
            (8, "Days sales outstanding", assumptions.flat_dso, "0.0"),
            (9, "Days inventory outstanding", assumptions.flat_dio, "0.0"),
            (10, "Days payable outstanding", assumptions.flat_dpo, "0.0"),
            (11, "Other current assets / revenue", assumptions.other_ca_pct_revenue, RATE),
            (12, "Accrued expenses / revenue", assumptions.accrued_pct_revenue, RATE),
            (13, "COGS / revenue", assumptions.costs["cogs"].flat_pct, RATE),
            (14, "Personnel / revenue", assumptions.costs["personnel"].flat_pct, RATE),
            (15, "R&D / revenue", assumptions.costs["rd"].flat_pct, RATE),
            (16, "Sales and marketing / revenue", assumptions.costs["sales_marketing"].flat_pct, RATE),
            (17, "G&A / revenue", assumptions.costs["ga"].flat_pct, RATE),
            (18, "Other operating cost / revenue", assumptions.costs["other_opex"].flat_pct, RATE),
            (19, "Minimum cash", assumptions.minimum_cash, MONEY),
            (20, "Terminal growth", assumptions.valuation.terminal_growth, RATE),
            (21, "Exit EBITDA multiple (cross-check)", assumptions.valuation.exit_ebitda_multiple, "0.00"),
            (22, "Diluted shares", assumptions.valuation.diluted_shares, "#,##0.00"),
            (23, "Operating cash requirement", assumptions.valuation.operating_cash_requirement, MONEY),
            (24, "Market value of equity", assumptions.wacc.market_value_equity, MONEY),
            (25, "Market value of debt", assumptions.wacc.market_value_debt, MONEY),
            (26, "Risk-free rate", assumptions.wacc.risk_free_rate, RATE),
            (27, "Beta", assumptions.wacc.beta, "0.00"),
            (28, "Equity risk premium", assumptions.wacc.equity_risk_premium, RATE),
            (29, "Pre-tax cost of debt", assumptions.wacc.pre_tax_cost_of_debt, RATE),
            (30, "Marginal tax rate", assumptions.wacc.marginal_tax_rate, RATE),
            (31, "Opening tax-loss balance", assumptions.opening_book_nol, MONEY),
            (32, "Opening unlevered tax-loss balance", assumptions.opening_unlevered_nol, MONEY),
            (33, "Remaining life of opening PP&E (years)", assumptions.existing_ppe_remaining_life, "0.00"),
            (34, "Useful life of new CapEx (years)", assumptions.new_ppe_life, "0.00"),
            (35, "Remaining life of opening intangibles (years)", assumptions.intangible_remaining_life, "0.00"),
            (36, "Minority interest", assumptions.valuation.minority_interest, MONEY),
            (37, "Preferred securities", assumptions.valuation.preferred_securities, MONEY),
            (38, "Other debt-like claims", assumptions.valuation.other_debt_like, MONEY),
            (39, "Non-operating assets", assumptions.valuation.non_operating_assets, MONEY),
        ]
        for row, label, value, fmt in rows:
            self._label("Assumptions", row, label)
            self._put("Assumptions", f"B{row}", _num(value), "input", fmt)
        self._label("Assumptions", 40, "Terminal method (gordon or exit_multiple)")
        self._put("Assumptions", "B40", assumptions.valuation.terminal_method, "input", "@")
        self._label("Assumptions", 42, "Discount convention (end or mid)")
        self._put("Assumptions", "B42", assumptions.valuation.discount_convention, "input", "@")
        self._label("Assumptions", 41, "Days in the period", True)
        self._label("Assumptions", 44, "Revenue growth rate", True)
        self._label("Assumptions", 45, "Scheduled debt repayment", True)
        for year in self.years:
            letter = self.col[year.period_id]
            self._put("Assumptions", f"{letter}41", year.days, "input", "0")
            if year.is_forecast:
                growth = assumptions.growth_rates[year.period_id]
                repayment = assumptions.repayments.get(year.period_id, Decimal(0))
                self._put("Assumptions", f"{letter}44", _num(growth), "input", RATE)
                self._put("Assumptions", f"{letter}45", _num(repayment), "input", MONEY)

    def _revenue(self) -> None:
        self._banner("Revenue", "Forecast revenue = prior revenue × (1 + that year's growth rate).")
        self._label("Revenue", 5, "Revenue")
        for year in self.years:
            letter = self.col[year.period_id]
            if year.is_forecast:
                prior = self.previous(year.period_id)
                self._put("Revenue", f"{letter}5", f"={self.col[prior]}5*(1+Assumptions!{letter}44)")
            else:
                self._put("Revenue", f"{letter}5", _num(year.income.get("revenue")), "input")

    def _costs(self) -> None:
        self._banner("OperatingCosts", "Line-item mode. Margin targets are not applied on this sheet.")
        lines = [
            (5, "Cost of goods sold", "cogs", 13),
            (6, "Personnel expenses", "personnel", 14),
            (7, "Research and development", "rd", 15),
            (8, "Sales and marketing", "sales_marketing", 16),
            (9, "General and administrative", "ga", 17),
            (10, "Other operating expenses", "other_opex", 18),
        ]
        for row, label, _key, assumption_row in lines:
            self._label("OperatingCosts", row, label)
        self._label("OperatingCosts", 11, "Cash operating costs")
        self._label("OperatingCosts", 12, "EBITDA")
        for year in self.years:
            letter = self.col[year.period_id]
            if year.is_forecast:
                for row, _label, _key, assumption_row in lines:
                    self._put("OperatingCosts", f"{letter}{row}", f"=Revenue!{letter}5*Assumptions!$B${assumption_row}")
            else:
                for row, _label, key, _assumption_row in lines:
                    self._put("OperatingCosts", f"{letter}{row}", _num(year.income.get(key)), "input")
            self._put("OperatingCosts", f"{letter}11", f"=SUM({letter}6:{letter}10)")
            self._put("OperatingCosts", f"{letter}12", f"=Revenue!{letter}5-{letter}5-{letter}11")
            self.wb["OperatingCosts"][f"{letter}12"].fill = TOTAL

    def _capex(self) -> None:
        self._banner("Capex", "Opening net book value is one straight-line layer. Each forecast year's CapEx is a new layer.")
        labels = {
            5: "Opening PP&E",
            6: "Capital expenditure",
            7: "Depreciation",
            8: "Disposals at net book value",
            9: "Impairments",
            10: "Closing PP&E",
            12: "Opening-layer balance",
            13: "Opening-layer depreciation",
            14: "Opening-layer closing balance",
            15: "Opening-layer annual charge",
            17: "Opening intangibles",
            18: "Amortisation",
            19: "Closing intangibles",
        }
        for row, label in labels.items():
            self._label("Capex", row, label)
        first = self.forecast_ids[0]
        first_letter = self.col[first]
        self._put("Capex", f"{first_letter}15", f"={first_letter}12/Assumptions!$B$33")
        for offset, period_id in enumerate(self.forecast_ids):
            base = 21 + offset * 3
            self._label("Capex", base, f"CapEx layer {period_id}")
            self._label("Capex", base + 1, f"Depreciation of {period_id} layer")
            self._label("Capex", base + 2, f"Remaining {period_id} layer")
        for period_id in self.forecast_ids:
            letter = self.col[period_id]
            prior = self.previous(period_id)
            prior_letter = self.col[prior]
            if prior == self.last_actual:
                self._put("Capex", f"{letter}5", f"=Balance!{prior_letter}9")
                self._put("Capex", f"{letter}12", f"=Balance!{prior_letter}9")
            else:
                # Total opening book value carries every surviving layer.
                # Row 12 continues only the original opening layer.
                self._put("Capex", f"{letter}5", f"={prior_letter}10")
                self._put("Capex", f"{letter}12", f"={prior_letter}14")
            self._put("Capex", f"{letter}13", f"=MIN({first_letter}15,{letter}12)")
            self._put("Capex", f"{letter}14", f"={letter}12-{letter}13")
            if letter != first_letter:
                self._put("Capex", f"{letter}15", f"={first_letter}15")
            self._put("Capex", f"{letter}6", f"=Revenue!{letter}5*Assumptions!$B$7")
            self._put("Capex", f"{letter}8", 0, "input")
            self._put("Capex", f"{letter}9", 0, "input")
            vintage_deps = []
            vintage_ends = []
            for offset, vintage in enumerate(self.forecast_ids):
                base = 21 + offset * 3
                vintage_letter = self.col[vintage]
                if self.ids.index(period_id) < self.ids.index(vintage):
                    self._put("Capex", f"{letter}{base}", 0)
                    self._put("Capex", f"{letter}{base + 1}", 0)
                    self._put("Capex", f"{letter}{base + 2}", 0)
                elif period_id == vintage:
                    self._put("Capex", f"{letter}{base}", f"={letter}6")
                    self._put("Capex", f"{letter}{base + 1}", f"={letter}{base}/Assumptions!$B$34")
                    self._put("Capex", f"{letter}{base + 2}", f"={letter}{base}-{letter}{base + 1}")
                else:
                    self._put("Capex", f"{letter}{base}", 0)
                    self._put(
                        "Capex",
                        f"{letter}{base + 1}",
                        f"=MIN({vintage_letter}{base}/Assumptions!$B$34,{prior_letter}{base + 2})",
                    )
                    self._put("Capex", f"{letter}{base + 2}", f"={prior_letter}{base + 2}-{letter}{base + 1}")
                vintage_deps.append(f"{letter}{base + 1}")
                vintage_ends.append(f"{letter}{base + 2}")
            self._put("Capex", f"{letter}7", "=" + "+".join([f"{letter}13", *vintage_deps]))
            self._put("Capex", f"{letter}10", f"={letter}5+{letter}6-{letter}7-{letter}8-{letter}9")
            if prior == self.last_actual:
                self._put("Capex", f"{letter}17", f"=Balance!{prior_letter}10")
            else:
                self._put("Capex", f"{letter}17", f"={prior_letter}19")
            self._put(
                "Capex",
                f"{letter}18",
                f"=MIN({self.col[first]}17/Assumptions!$B$35,{letter}17)",
            )
            self._put("Capex", f"{letter}19", f"={letter}17-{letter}18")
        # Historical capex is a fact from the cash flow statement.
        self._label("Capex", 4, "")
        for year in self.years:
            if year.is_forecast:
                continue
            letter = self.col[year.period_id]
            capex = None if not year.cash_flow else year.cash_flow.get("capex")
            self._put("Capex", f"{letter}6", _num(capex), "input")
            self._put("Capex", f"{letter}7", _num(year.income.get("depreciation")), "input")
            self._put("Capex", f"{letter}18", _num(year.income.get("amortisation")), "input")

    def _debt(self) -> None:
        self._banner("Debt", "Interest uses the opening balance, so a same-year revolver draw does not change this year's interest.")
        for row, label in (
            (5, "Opening short-term debt"),
            (6, "Opening long-term debt"),
            (7, "Opening revolver"),
            (8, "New borrowings"),
            (9, "Repayment of short-term debt"),
            (10, "Repayment of long-term debt"),
            (11, "Revolver draw / (repayment)"),
            (12, "Closing short-term debt"),
            (13, "Closing long-term debt"),
            (14, "Closing revolver"),
            (15, "Interest expense"),
            (16, "Interest income"),
        ):
            self._label("Debt", row, label)
        for period_id in self.forecast_ids:
            letter = self.col[period_id]
            prior = self.col[self.previous(period_id)]
            self._put("Debt", f"{letter}5", f"=Balance!{prior}17")
            self._put("Debt", f"{letter}6", f"=Balance!{prior}19")
            self._put("Debt", f"{letter}7", f"=Balance!{prior}18")
            self._put("Debt", f"{letter}8", 0, "input")
            self._put("Debt", f"{letter}9", f"=MIN(Assumptions!{letter}45,{letter}5+{letter}8)")
            self._put("Debt", f"{letter}10", f"=Assumptions!{letter}45-{letter}9")
            self._put("Debt", f"{letter}15", f"=Assumptions!$B$4*({letter}5+{letter}6+{letter}7)")
            self._put("Debt", f"{letter}16", f"=Assumptions!$B$5*Balance!{prior}5")
            # The revolver flow is finished after the cash flow sheet exists. A placeholder
            # is replaced in _cash_flow once the cash-before-revolver row is known.
            self._put("Debt", f"{letter}12", f"={letter}5+{letter}8-{letter}9")
            self._put("Debt", f"{letter}13", f"={letter}6-{letter}10")
            self._put("Debt", f"{letter}14", f"={letter}7+{letter}11")

    def _working_capital(self) -> None:
        self._banner("WorkingCapital", "Forecast balances are period-end: flow × days ratio / days in the period. Cash and debt are excluded.")
        for row, label in (
            (5, "Accounts receivable"),
            (6, "Inventory"),
            (7, "Other current assets"),
            (8, "Accounts payable"),
            (9, "Accrued expenses"),
            (10, "Other operating current liabilities"),
            (11, "Operating net working capital"),
            (12, "Increase in operating net working capital"),
        ):
            self._label("WorkingCapital", row, label)
        for year in self.years:
            letter = self.col[year.period_id]
            if year.is_forecast:
                self._put("WorkingCapital", f"{letter}5", f"=Revenue!{letter}5*Assumptions!$B$8/Assumptions!{letter}41")
                self._put("WorkingCapital", f"{letter}6", f"=OperatingCosts!{letter}5*Assumptions!$B$9/Assumptions!{letter}41")
                self._put("WorkingCapital", f"{letter}7", f"=Revenue!{letter}5*Assumptions!$B$11")
                self._put("WorkingCapital", f"{letter}8", f"=OperatingCosts!{letter}5*Assumptions!$B$10/Assumptions!{letter}41")
                self._put("WorkingCapital", f"{letter}9", f"=Revenue!{letter}5*Assumptions!$B$12")
                self._put("WorkingCapital", f"{letter}10", f"=Revenue!{letter}5*0")
            else:
                balance = year.balance
                for row, key in (
                    (5, "accounts_receivable"),
                    (6, "inventory"),
                    (7, "other_current_assets"),
                    (8, "accounts_payable"),
                    (9, "accrued_expenses"),
                    (10, "other_operating_cl"),
                ):
                    self._put("WorkingCapital", f"{letter}{row}", _num(balance.get(key)), "input")
            self._put("WorkingCapital", f"{letter}11", f"={letter}5+{letter}6+{letter}7-{letter}8-{letter}9-{letter}10")
            prior = self.previous(year.period_id)
            if prior is not None and year.is_forecast:
                self._put("WorkingCapital", f"{letter}12", f"={letter}11-{self.col[prior]}11")

    def _income(self) -> None:
        self._banner("Income", "Calculated lines are formulas. Forecast operations come from Revenue, OperatingCosts, Capex and Debt.")
        labels = (
            (5, "Revenue"),
            (6, "Cost of goods sold"),
            (7, "Gross profit"),
            (8, "Personnel expenses"),
            (9, "Research and development"),
            (10, "Sales and marketing"),
            (11, "General and administrative"),
            (12, "Other operating expenses"),
            (13, "EBITDA"),
            (14, "Depreciation"),
            (15, "Amortisation"),
            (16, "Impairments"),
            (17, "EBIT"),
            (18, "Interest income"),
            (19, "Interest expense"),
            (20, "Other non-operating items"),
            (21, "Profit before tax"),
            (22, "Income tax expense"),
            (23, "Net income"),
            (25, "Opening book tax-loss balance"),
            (26, "Closing book tax-loss balance"),
            (27, "Unlevered tax"),
            (28, "Opening unlevered tax-loss balance"),
            (29, "Closing unlevered tax-loss balance"),
        )
        for row, label in labels:
            self._label("Income", row, label)
        first = self.forecast_ids[0]
        for year in self.years:
            letter = self.col[year.period_id]
            self._put("Income", f"{letter}5", f"=Revenue!{letter}5")
            self._put("Income", f"{letter}6", f"=OperatingCosts!{letter}5")
            self._put("Income", f"{letter}7", f"={letter}5-{letter}6")
            self._put("Income", f"{letter}8", f"=OperatingCosts!{letter}6")
            self._put("Income", f"{letter}9", f"=OperatingCosts!{letter}7")
            self._put("Income", f"{letter}10", f"=OperatingCosts!{letter}8")
            self._put("Income", f"{letter}11", f"=OperatingCosts!{letter}9")
            self._put("Income", f"{letter}12", f"=OperatingCosts!{letter}10")
            self._put("Income", f"{letter}13", f"={letter}7-SUM({letter}8:{letter}12)")
            if year.is_forecast:
                self._put("Income", f"{letter}14", f"=Capex!{letter}7")
                self._put("Income", f"{letter}15", f"=Capex!{letter}18")
                self._put("Income", f"{letter}16", f"=Capex!{letter}9")
                self._put("Income", f"{letter}18", f"=Debt!{letter}16")
                self._put("Income", f"{letter}19", f"=Debt!{letter}15")
                self._put("Income", f"{letter}20", 0, "input")
                prior = self.previous(year.period_id)
                if year.period_id == first:
                    self._put("Income", f"{letter}25", "=Assumptions!$B$31")
                    self._put("Income", f"{letter}28", "=Assumptions!$B$32")
                else:
                    self._put("Income", f"{letter}25", f"={self.col[prior]}26")
                    self._put("Income", f"{letter}28", f"={self.col[prior]}29")
                self._put(
                    "Income",
                    f"{letter}22",
                    f"=(MAX(0,{letter}21)-MIN({letter}25,MAX(0,{letter}21)))*Assumptions!$B$3",
                )
                self._put(
                    "Income",
                    f"{letter}26",
                    f"={letter}25-MIN({letter}25,MAX(0,{letter}21))+MAX(-{letter}21,0)",
                )
                self._put(
                    "Income",
                    f"{letter}27",
                    f"=(MAX(0,{letter}17)-MIN({letter}28,MAX(0,{letter}17)))*Assumptions!$B$3",
                )
                self._put(
                    "Income",
                    f"{letter}29",
                    f"={letter}28-MIN({letter}28,MAX(0,{letter}17))+MAX(-{letter}17,0)",
                )
            else:
                for row, key in (
                    (14, "depreciation"),
                    (15, "amortisation"),
                    (16, "impairment"),
                    (18, "interest_income"),
                    (19, "interest_expense"),
                    (20, "other_non_operating"),
                    (22, "tax_expense"),
                ):
                    self._put("Income", f"{letter}{row}", _num(year.income.get(key)), "input")
            self._put("Income", f"{letter}17", f"={letter}13-{letter}14-{letter}15-{letter}16")
            self._put("Income", f"{letter}21", f"={letter}17+{letter}18-{letter}19+{letter}20")
            if not year.is_forecast:
                self._put("Income", f"{letter}23", f"={letter}21-{letter}22")
            else:
                self._put("Income", f"{letter}23", f"={letter}21-{letter}22")
            self.wb["Income"][f"{letter}23"].fill = TOTAL

    def _cash_flow(self) -> None:
        self._banner("CashFlow", "Indirect method. Closing cash is this statement's result and is the balance-sheet cash balance.")
        labels = (
            (5, "Net income"),
            (6, "Depreciation, amortisation and impairments"),
            (7, "Gain on disposal"),
            (8, "Other non-cash adjustments"),
            (9, "Increase in operating net working capital"),
            (10, "Cash flow from operations"),
            (12, "Capital expenditure"),
            (13, "Disposal proceeds"),
            (14, "Cash flow from investing"),
            (16, "New borrowings"),
            (17, "Principal repayments"),
            (18, "Revolver draw / (repayment)"),
            (19, "Dividends"),
            (20, "Cash flow from financing"),
            (22, "Net change in cash"),
            (23, "Opening cash"),
            (24, "Closing cash"),
            (26, "Cash before the revolver"),
        )
        for row, label in labels:
            self._label("CashFlow", row, label)
        for year in self.years:
            letter = self.col[year.period_id]
            if not year.is_forecast:
                flow = year.cash_flow or {}
                for row, key in (
                    (5, "net_income"),
                    (6, "depreciation_amortisation"),
                    (7, "disposal_gain"),
                    (8, "other_non_cash"),
                    (9, "change_in_nwc"),
                    (10, "cash_from_operations"),
                    (12, "capex"),
                    (13, "disposal_proceeds"),
                    (14, "cash_from_investing"),
                    (16, "borrowings"),
                    (17, "repayments"),
                    (18, "revolver_draw"),
                    (19, "dividends"),
                    (20, "cash_from_financing"),
                    (22, "net_change_in_cash"),
                    (23, "opening_cash"),
                    (24, "closing_cash"),
                ):
                    self._put("CashFlow", f"{letter}{row}", _num(flow.get(key)), "input")
                continue
            prior = self.col[self.previous(year.period_id)]
            self._put("CashFlow", f"{letter}5", f"=Income!{letter}23")
            self._put("CashFlow", f"{letter}6", f"=Income!{letter}14+Income!{letter}15+Income!{letter}16")
            self._put("CashFlow", f"{letter}7", f"={letter}13-Capex!{letter}8")
            self._put("CashFlow", f"{letter}8", 0, "input")
            self._put("CashFlow", f"{letter}9", f"=WorkingCapital!{letter}12")
            self._put("CashFlow", f"{letter}10", f"={letter}5+{letter}6-{letter}7+{letter}8-{letter}9")
            self._put("CashFlow", f"{letter}12", f"=Capex!{letter}6")
            self._put("CashFlow", f"{letter}13", 0, "input")
            self._put("CashFlow", f"{letter}14", f"=-{letter}12+{letter}13")
            self._put("CashFlow", f"{letter}16", f"=Debt!{letter}8")
            self._put("CashFlow", f"{letter}17", f"=Debt!{letter}9+Debt!{letter}10")
            self._put("CashFlow", f"{letter}19", f"=MAX(0,Income!{letter}23)*Assumptions!$B$6")
            self._put("CashFlow", f"{letter}23", f"=Balance!{prior}5")
            self._put(
                "CashFlow",
                f"{letter}26",
                f"={letter}23+{letter}10+{letter}14+{letter}16-{letter}17-{letter}19",
            )
            self._put(
                "Debt",
                f"{letter}11",
                f'=IF(CashFlow!{letter}26-Debt!{letter}7>=Assumptions!$B$19,-Debt!{letter}7,Assumptions!$B$19-CashFlow!{letter}26)',
            )
            self._put("CashFlow", f"{letter}18", f"=Debt!{letter}11")
            self._put("CashFlow", f"{letter}20", f"={letter}16-{letter}17+{letter}18-{letter}19")
            self._put("CashFlow", f"{letter}22", f"={letter}10+{letter}14+{letter}20")
            self._put("CashFlow", f"{letter}24", f"={letter}23+{letter}22")

    def _balance(self) -> None:
        self._banner("Balance", "Forecast cash is the cash flow statement's closing cash. There is no balancing plug.")
        labels = (
            (5, "Cash and cash equivalents"),
            (6, "Accounts receivable"),
            (7, "Inventory"),
            (8, "Other current assets"),
            (9, "Property, plant and equipment"),
            (10, "Intangible assets"),
            (11, "Other non-current assets"),
            (12, "Total assets"),
            (14, "Accounts payable"),
            (15, "Accrued expenses"),
            (16, "Other operating current liabilities"),
            (17, "Current debt"),
            (18, "Minimum-cash revolver"),
            (19, "Long-term debt"),
            (20, "Other liabilities"),
            (21, "Total liabilities"),
            (23, "Share capital"),
            (24, "Additional paid-in capital"),
            (25, "Retained earnings"),
            (26, "Other equity"),
            (27, "Total equity"),
            (28, "Total liabilities and equity"),
        )
        for row, label in labels:
            self._label("Balance", row, label)
        for year in self.years:
            letter = self.col[year.period_id]
            if not year.is_forecast:
                balance = year.balance
                for row, key in (
                    (5, "cash"),
                    (6, "accounts_receivable"),
                    (7, "inventory"),
                    (8, "other_current_assets"),
                    (9, "ppe_net"),
                    (10, "intangibles"),
                    (11, "other_non_current_assets"),
                    (14, "accounts_payable"),
                    (15, "accrued_expenses"),
                    (16, "other_operating_cl"),
                    (17, "short_term_debt"),
                    (18, "revolver"),
                    (19, "long_term_debt"),
                    (20, "other_liabilities"),
                    (23, "share_capital"),
                    (24, "apic"),
                    (25, "retained_earnings"),
                    (26, "other_equity"),
                ):
                    self._put("Balance", f"{letter}{row}", _num(balance.get(key)), "input")
            else:
                prior = self.col[self.previous(year.period_id)]
                self._put("Balance", f"{letter}5", f"=CashFlow!{letter}24")
                self._put("Balance", f"{letter}6", f"=WorkingCapital!{letter}5")
                self._put("Balance", f"{letter}7", f"=WorkingCapital!{letter}6")
                self._put("Balance", f"{letter}8", f"=WorkingCapital!{letter}7")
                self._put("Balance", f"{letter}9", f"=Capex!{letter}10")
                self._put("Balance", f"{letter}10", f"=Capex!{letter}19")
                self._put("Balance", f"{letter}11", f"={prior}11")
                self._put("Balance", f"{letter}14", f"=WorkingCapital!{letter}8")
                self._put("Balance", f"{letter}15", f"=WorkingCapital!{letter}9")
                self._put("Balance", f"{letter}16", f"=WorkingCapital!{letter}10")
                self._put("Balance", f"{letter}17", f"=Debt!{letter}12")
                self._put("Balance", f"{letter}18", f"=Debt!{letter}14")
                self._put("Balance", f"{letter}19", f"=Debt!{letter}13")
                self._put("Balance", f"{letter}20", f"={prior}20")
                self._put("Balance", f"{letter}23", f"={prior}23")
                self._put("Balance", f"{letter}24", f"={prior}24")
                self._put("Balance", f"{letter}25", f"={prior}25+Income!{letter}23-CashFlow!{letter}19")
                self._put("Balance", f"{letter}26", f"={prior}26")
            self._put("Balance", f"{letter}12", f"=SUM({letter}5:{letter}11)")
            self._put("Balance", f"{letter}21", f"={letter}14+{letter}15+{letter}16+{letter}17+{letter}18+{letter}19+{letter}20")
            self._put("Balance", f"{letter}27", f"=SUM({letter}23:{letter}26)")
            self._put("Balance", f"{letter}28", f"={letter}21+{letter}27")
            self.wb["Balance"][f"{letter}12"].fill = TOTAL
            self.wb["Balance"][f"{letter}28"].fill = TOTAL

    def _wacc(self) -> None:
        self._banner("WACC", "CAPM cost of equity. Market values are the inputs in Assumptions, not book equity.")
        rows = (
            (5, "Risk-free rate", "=Assumptions!B26"),
            (6, "Beta", "=Assumptions!B27"),
            (7, "Equity risk premium", "=Assumptions!B28"),
            (8, "Cost of equity", "=B5+B6*B7"),
            (9, "Pre-tax cost of debt", "=Assumptions!B29"),
            (10, "Marginal tax rate", "=Assumptions!B30"),
            (11, "After-tax cost of debt", "=B9*(1-B10)"),
            (12, "Market value of equity", "=Assumptions!B24"),
            (13, "Market value of debt", "=Assumptions!B25"),
            (14, "Equity weight", "=B12/(B12+B13)"),
            (15, "Debt weight", "=B13/(B12+B13)"),
            (16, "WACC", "=B14*B8+B15*B11"),
        )
        for row, label, formula in rows:
            self._label("WACC", row, label)
            fmt = RATE if row in (5, 7, 8, 9, 10, 11, 14, 15, 16) else ("0.00" if row == 6 else MONEY)
            self._put("WACC", f"B{row}", formula, fmt=fmt)
        self.wb["WACC"]["B16"].fill = TOTAL

    def _dcf(self) -> None:
        self._banner(
            "DCF",
            "Discount exponents follow Assumptions!B42. Terminal value follows Assumptions!B40. The other terminal method is the cross-check.",
        )
        for row, label in (
            (5, "EBIT"),
            (6, "Unlevered tax"),
            (7, "NOPAT"),
            (8, "Depreciation and amortisation"),
            (9, "Capital expenditure"),
            (10, "Increase in operating net working capital"),
            (11, "Unlevered free cash flow"),
            (12, "Discount exponent"),
            (13, "Discount factor"),
            (14, "Present value of unlevered free cash flow"),
        ):
            self._label("DCF", row, label)
        self._label("DCF", 16, "Terminal value")
        self._label("DCF", 17, "Present value of terminal value")
        self._label("DCF", 18, "Enterprise value")
        self._label("DCF", 19, "Exit-multiple terminal value (cross-check)")
        for offset, period_id in enumerate(self.forecast_ids, start=1):
            letter = self.col[period_id]
            self._put("DCF", f"{letter}5", f"=Income!{letter}17")
            self._put("DCF", f"{letter}6", f"=Income!{letter}27")
            self._put("DCF", f"{letter}7", f"={letter}5-{letter}6")
            self._put("DCF", f"{letter}8", f"=Income!{letter}14+Income!{letter}15+Income!{letter}16")
            self._put("DCF", f"{letter}9", f"=Capex!{letter}6")
            self._put("DCF", f"{letter}10", f"=WorkingCapital!{letter}12")
            self._put("DCF", f"{letter}11", f"={letter}7+{letter}8-{letter}9+CashFlow!{letter}13-{letter}10")
            self._put("DCF", f"{letter}12", f'=IF(Assumptions!$B$42="mid",{offset}-0.5,{offset})', fmt="0.00")
            self._put("DCF", f"{letter}13", f"=1/(1+WACC!$B$16)^{letter}12")
            self._put("DCF", f"{letter}14", f"={letter}11*{letter}13")
        last = self.col[self.forecast_ids[-1]]
        first = self.col[self.forecast_ids[0]]
        gordon = f'IF(WACC!B16<=Assumptions!B20,"n/m",{last}11*(1+Assumptions!B20)/(WACC!B16-Assumptions!B20))'
        exit_value = f"Income!{last}13*Assumptions!B21"
        self._put("DCF", "B16", f'=IF(Assumptions!B40="exit_multiple",{exit_value},{gordon})')
        self._put("DCF", "B17", f'=IF(B16="n/m","n/m",B16*{last}13)')
        self._put("DCF", "B18", f'=IF(B17="n/m","n/m",SUM({first}14:{last}14)+B17)')
        self._put("DCF", "B19", f'=IF(Assumptions!B40="exit_multiple",{gordon},{exit_value})')
        self._put("DCF", "B20", f'=IF(B19="n/m","n/m",SUM({first}14:{last}14)+B19*{last}13)')
        self._label("DCF", 20, "Cross-check enterprise value")
        self.wb["DCF"]["B18"].fill = TOTAL

    def _summary(self) -> None:
        ws = self.wb["Summary"]
        ws["A1"] = self.model.name
        ws["A1"].font = Font(name="Calibri", size=20, bold=True, color=NAVY)
        ws["A2"] = self.model.synthetic_notice if self.model.synthetic else "Company forecast"
        ws["A2"].font = Font(name="Calibri", size=10, color="8A5A00")
        ws.column_dimensions["A"].width = 52
        ws.column_dimensions["B"].width = 22
        items = (
            (4, "Currency", self.model.currency),
            (5, "Units", "Absolute currency units"),
            (6, "Enterprise value", "=DCF!B18"),
            (7, "Equity value", "=B6-Balance!{c}17-Balance!{c}18-Balance!{c}19-Assumptions!B38-Assumptions!B36-Assumptions!B37+MAX(0,Balance!{c}5-Assumptions!B23)+Assumptions!B39"),
            (8, "Intrinsic value per share", '=IF(Assumptions!B22<=0,"n/m",B7/Assumptions!B22)'),
            (9, "WACC", "=WACC!B16"),
            (10, "Terminal growth", "=Assumptions!B20"),
            (11, "Present value of explicit UFCF", "=SUM(DCF!{f}14:DCF!{l}14)"),
            (12, "Present value of terminal value", "=DCF!B17"),
            (13, "Terminal value share of enterprise value", "=IF(B6=0,\"n/m\",B12/B6)"),
            (14, "Valuation date", str(self.model.assumptions.valuation.valuation_date)),
            (15, "Scenario", self.model.assumptions.name),
            (16, "Discount convention", "=Assumptions!B42"),
        )
        last_actual = self.col[self.last_actual]
        first = self.col[self.forecast_ids[0]]
        last = self.col[self.forecast_ids[-1]]
        for row, label, value in items:
            self._label("Summary", row, label)
            if isinstance(value, str) and value.startswith("="):
                formula = value.format(c=last_actual, f=first, l=last)
                if row == 16:
                    fmt = "@"
                elif row in (9, 10, 13):
                    fmt = RATE
                else:
                    fmt = MONEY
                self._put("Summary", f"B{row}", formula, fmt=fmt)
            else:
                self._put("Summary", f"B{row}", value, "calc", fmt="@")
        ws["B6"].fill = TOTAL
        assumptions = self.model.assumptions
        supported = (
            assumptions.revenue_method == "growth"
            and assumptions.cost_mode == "line_item"
            and assumptions.capex_method == "pct_revenue"
            and assumptions.interest_basis == "opening"
        )
        if supported:
            ws["A3"] = "The live case is the base case on the calculation sheets. Bull and bear cases are on Scenarios."
        else:
            ws["A3"] = (
                "The formula sheets use revenue growth, percentage-of-revenue costs, percentage-of-revenue capex and opening-balance interest. "
                f"This model is revenue {assumptions.revenue_method}, costs {assumptions.cost_mode}, "
                f"capex {assumptions.capex_method}, interest {assumptions.interest_basis}. "
                "Those formula sheets will not reproduce this case."
            )
        ws["A3"].font = Font(name="Calibri", size=10, italic=True, color="52606D")

    def _historical_links(self) -> None:
        self._banner("Historical", "Historical columns link back to the statement sheets. They are not a second set of numbers.")
        self._label("Historical", 5, "Revenue")
        self._label("Historical", 6, "EBITDA")
        self._label("Historical", 7, "Net income")
        self._label("Historical", 8, "Cash")
        for year in self.years:
            if year.is_forecast:
                continue
            letter = self.col[year.period_id]
            self._put("Historical", f"{letter}5", f"=Revenue!{letter}5")
            self._put("Historical", f"{letter}6", f"=Income!{letter}13")
            self._put("Historical", f"{letter}7", f"=Income!{letter}23")
            self._put("Historical", f"{letter}8", f"=Balance!{letter}5")

    def _checks(self) -> None:
        self._banner("Checks", "A zero difference means the formula sheets still satisfy the accounting identity.")
        self._label("Checks", 5, "Assets minus liabilities and equity")
        self._label("Checks", 6, "Balance-sheet cash minus cash-flow cash")
        self._label("Checks", 7, "PP&E identity")
        for year in self.years:
            letter = self.col[year.period_id]
            self._put("Checks", f"{letter}5", f"=Balance!{letter}12-Balance!{letter}28")
            if year.is_forecast:
                self._put("Checks", f"{letter}6", f"=Balance!{letter}5-CashFlow!{letter}24")
                self._put("Checks", f"{letter}7", f"=Capex!{letter}10-(Capex!{letter}5+Capex!{letter}6-Capex!{letter}7-Capex!{letter}8-Capex!{letter}9)")

    def _sensitivities(self) -> None:
        self._banner("Sensitivities", "The WACC and growth grid is formula-driven from the UFCF row. Operating grids are full-model engine results.")
        ws = self.wb["Sensitivities"]
        ws["A5"] = "WACC \\ terminal growth"
        ws["A5"].font = Font(name="Calibri", size=10, bold=True)
        growth_rates = [Decimal("-0.005"), Decimal("0"), Decimal("0.005"), Decimal("0.01"), Decimal("0.015")]
        wacc_offsets = [Decimal("-0.01"), Decimal("-0.005"), Decimal("0"), Decimal("0.005"), Decimal("0.01")]
        last = self.col[self.forecast_ids[-1]]
        for index, offset in enumerate(growth_rates):
            letter = get_column_letter(3 + index)
            if offset == 0:
                self._put("Sensitivities", f"{letter}6", "=Assumptions!B20", fmt=RATE)
            else:
                sign = "+" if offset > 0 else ""
                self._put("Sensitivities", f"{letter}6", f"=Assumptions!B20{sign}{offset}", fmt=RATE)
        invalid_letter = get_column_letter(3 + len(growth_rates))
        self._put("Sensitivities", f"{invalid_letter}6", f"=WACC!B16-0.01", fmt=RATE)
        ws[f"{invalid_letter}5"] = "Invalid boundary"
        for row_offset, wacc_offset in enumerate(wacc_offsets):
            row = 7 + row_offset
            if wacc_offset == 0:
                self._put("Sensitivities", f"B{row}", "=WACC!B16", fmt=RATE)
            else:
                sign = "+" if wacc_offset > 0 else ""
                self._put("Sensitivities", f"B{row}", f"=WACC!B16{sign}{wacc_offset}", fmt=RATE)
            for index in range(len(growth_rates) + 1):
                letter = get_column_letter(3 + index)
                parts = []
                for period_id in self.forecast_ids:
                    column = self.col[period_id]
                    parts.append(f"DCF!{column}11/(1+$B{row})^DCF!{column}12")
                terminal = f"DCF!{last}11*(1+{letter}$6)/($B{row}-{letter}$6)/(1+$B{row})^DCF!{last}12"
                formula = f'=IF($B{row}<={letter}$6,"n/m",' + "+".join(parts) + f"+{terminal})"
                self._put("Sensitivities", f"{letter}{row}", formula)
        ws["A13"] = "WACC versus exit multiple. Each cell discounts the same unlevered free cash flow at that WACC and exit multiple."
        ws["A13"].font = Font(name="Calibri", size=10, italic=True, color="52606D")
        multiple_offsets = [Decimal("-1"), Decimal("-0.5"), Decimal("0"), Decimal("0.5"), Decimal("1")]
        for index, offset in enumerate(multiple_offsets):
            letter = get_column_letter(3 + index)
            if offset == 0:
                self._put("Sensitivities", f"{letter}15", "=Assumptions!B21", fmt="0.00")
            else:
                sign = "+" if offset > 0 else ""
                self._put("Sensitivities", f"{letter}15", f"=Assumptions!B21{sign}{offset}", fmt="0.00")
        for row_offset, wacc_offset in enumerate(wacc_offsets):
            row = 16 + row_offset
            if wacc_offset == 0:
                self._put("Sensitivities", f"B{row}", "=WACC!B16", fmt=RATE)
            else:
                sign = "+" if wacc_offset > 0 else ""
                self._put("Sensitivities", f"B{row}", f"=WACC!B16{sign}{wacc_offset}", fmt=RATE)
            for index in range(len(multiple_offsets)):
                letter = get_column_letter(3 + index)
                parts = [
                    f"DCF!{self.col[period_id]}11/(1+$B{row})^DCF!{self.col[period_id]}12"
                    for period_id in self.forecast_ids
                ]
                terminal = f"Income!{last}13*{letter}$15/(1+$B{row})^DCF!{last}12"
                formula = f'=IF({letter}$15<=0,"n/m",' + "+".join(parts) + f"+{terminal})"
                self._put("Sensitivities", f"{letter}{row}", formula)
        ws["A22"] = (
            "Operating sensitivity tables below are values produced by rebuilding the three-statement model in the calculation engine. "
            "They are not a percentage applied to enterprise value. Excel cannot hold a second full statement model in every cell of those grids."
        )
        ws["A22"].font = Font(name="Calibri", size=10, italic=True, color="52606D")
        row_cursor = 24
        for table in self.result.sensitivities:
            if table.name in {"WACC versus terminal growth", "WACC versus exit multiple"}:
                continue
            self._label("Sensitivities", row_cursor, table.name, True)
            row_cursor += 1
            self._put("Sensitivities", f"A{row_cursor}", table.row_label + " \\ " + table.column_label, "section")
            for index, key in enumerate(table.column_keys):
                self._put("Sensitivities", f"{get_column_letter(3 + index)}{row_cursor}", float(Decimal(key)), "calc", "0.00%")
            for r_index, key in enumerate(table.row_keys):
                row_cursor += 1
                self._put("Sensitivities", f"B{row_cursor}", float(Decimal(key)), "calc", "0.00%")
                for c_index, value in enumerate(table.values[r_index]):
                    letter = get_column_letter(3 + c_index)
                    if table.statuses[r_index][c_index] != "ok" or value is None:
                        self._put("Sensitivities", f"{letter}{row_cursor}", "n/m", "calc", "@")
                    else:
                        self._put("Sensitivities", f"{letter}{row_cursor}", float(Decimal(value)), "calc")
            row_cursor += 2

    def _scenarios(self) -> None:
        self._banner("Scenarios", "Bull and bear revenue paths are formulas. Full valuations are engine recalculations from those assumption sets.")
        ws = self.wb["Scenarios"]
        ws["A5"] = "Case"
        ws["B5"] = "Base"
        ws["C5"] = "Bull"
        ws["D5"] = "Bear"
        for cell in ("B5", "C5", "D5"):
            ws[cell].font = WHITE
            ws[cell].fill = HEADER
        ws["A6"] = "Revenue growth offset versus base"
        self._put("Scenarios", "B6", 0, "input", RATE)
        self._put("Scenarios", "C6", 0.02, "input", RATE)
        self._put("Scenarios", "D6", -0.02, "input", RATE)
        ws["A7"] = "COGS percentage"
        self._put("Scenarios", "B7", "=Assumptions!B13", fmt=RATE)
        self._put("Scenarios", "C7", "=Assumptions!B13-0.015", fmt=RATE)
        self._put("Scenarios", "D7", "=Assumptions!B13+0.015", fmt=RATE)
        ws["A9"] = "Forecast revenue"
        last_actual = self.col[self.last_actual]
        start_row = 10
        for index, period_id in enumerate(self.forecast_ids):
            row = start_row + index
            ws[f"A{row}"] = period_id
            for column, case in (("B", "B"), ("C", "C"), ("D", "D")):
                if index == 0:
                    formula = f"=Revenue!{last_actual}5*(1+Assumptions!{self.col[period_id]}44+{case}$6)"
                else:
                    formula = f"={column}{row - 1}*(1+Assumptions!{self.col[period_id]}44+{case}$6)"
                self._put("Scenarios", f"{column}{row}", formula)
        value_row = start_row + len(self.forecast_ids) + 2
        ws[f"A{value_row}"] = "Enterprise value from the full model"
        ws[f"A{value_row}"].font = Font(name="Calibri", size=10, italic=True)
        base_ev = self.result.dcf.enterprise_value if self.result.dcf else None
        bull = self.result.scenarios.get("Bull")
        bear = self.result.scenarios.get("Bear")
        self._put("Scenarios", f"B{value_row}", _num(base_ev), "calc")
        self._put("Scenarios", f"C{value_row}", _num(bull.dcf.enterprise_value if bull and bull.dcf else None), "calc")
        self._put("Scenarios", f"D{value_row}", _num(bear.dcf.enterprise_value if bear and bear.dcf else None), "calc")
        note = value_row + 2
        ws[f"A{note}"] = (
            "Enterprise values in the row above are outputs of the ValuationLab calculation engine, "
            "rebuilt from each case's assumptions. They are not a haircut to the base value. "
            "The base-case calculation sheets remain the formula-linked model."
        )
        ws[f"A{note}"].font = Font(name="Calibri", size=10, color="52606D")

    def _sources(self) -> None:
        ws = self.wb["Sources"]
        ws["A1"] = "Sources and modelling notes"
        ws["A1"].font = Font(name="Calibri", size=16, bold=True, color=NAVY)
        notes = [
            self.model.synthetic_notice,
            "Currency: " + self.model.currency + ". Amounts are absolute units, not thousands or millions.",
            "Blue font and yellow shading mark inputs. Green font marks a formula that reads another sheet. Black font marks a formula on the same sheet.",
            "Interest is calculated on opening interest-bearing debt, including the opening revolver.",
            "The minimum-cash revolver is explicit. It is solved from cash before the revolver because opening-balance interest does not depend on the current draw.",
            "Forecast working-capital balances use the period-end convention. Historical days ratios shown in the application may use averages and are not the forecast drivers.",
            "Unlevered free cash flow does not deduct interest. Tax in that bridge is tax on EBIT after the unlevered loss balance.",
            "Enterprise value uses the terminal method in Assumptions!B40. The other method is the cross-check on the DCF sheet.",
            "Discount exponents use Assumptions!B42. Mid-year exponents are the year number minus one half.",
            "Equity value subtracts valuation-date interest-bearing debt and adds cash above the operating-cash requirement. Forecast cash is not added again.",
            "Market-value inputs are manual. This workbook does not contain a live price.",
            "WACC inputs: " + (self.model.assumptions.wacc.risk_free_source if self.model.assumptions.wacc else ""),
            self.model.assumptions.notes,
        ]
        for index, note in enumerate(notes, start=3):
            ws[f"A{index}"] = note
            ws[f"A{index}"].font = Font(name="Calibri", size=10)
            ws.row_dimensions[index].height = 32
        ws.column_dimensions["A"].width = 140


def build_workbook(model: CompanyModel, result: ModelResult) -> bytes:
    return ModelWorkbook(model, result).build()
